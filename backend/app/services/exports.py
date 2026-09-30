import io
from datetime import datetime
from typing import Any

from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from app.models import CampaignMode, IdentityMode


def export_poll_excel(campaign, result: dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Poll Summary"
    ws.append(["Campaign", campaign.name])
    ws.append(["Mode", "Poll + Comment"])
    ws.append(["Target", campaign.target_user.full_name if campaign.target_user else ""])
    ws.append(["Identity mode", campaign.identity_mode.value])
    ws.append(["Average score", result.get("average_score")])
    ws.append(["Response rate %", result.get("response_rate")])
    ws.append([])
    ws.append(["Score", "Votes"])
    for score, count in sorted((result.get("distribution") or {}).items()):
        ws.append([score, count])
    ws2 = wb.create_sheet("Comments")
    ws2.append(["Score", "Comment", "Respondent"])
    for c in result.get("comments") or []:
        name = c.get("respondent_name") or ""
        if campaign.identity_mode != IdentityMode.NAMED:
            name = ""
        ws2.append([c.get("score"), c.get("comment") or "", name])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_grade_excel(campaign, result: dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Grade Ranking"
    ws.append(["Campaign", campaign.name])
    ws.append(["Group Head", campaign.group_head.full_name if campaign.group_head else ""])
    ws.append(["Rank direction", campaign.rank_direction.value])
    ws.append(["Final", result.get("is_final")])
    ws.append([])
    ws.append(["Rank", "Employee", "Employee ID"])
    for row in result.get("ranking") or []:
        ws.append([row.get("rank"), row.get("employee_name"), row.get("employee_id")])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_poll_pdf(campaign, result: dict[str, Any]) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    width, height = A4
    y = height - 50
    c.setFont("Helvetica-Bold", 14)
    c.drawString(40, y, "Poll Campaign Report")
    y -= 24
    c.setFont("Helvetica", 11)
    lines = [
        f"Campaign: {campaign.name}",
        f"Target: {campaign.target_user.full_name if campaign.target_user else '-'}",
        f"Identity: {campaign.identity_mode.value}",
        f"Invited: {result.get('invited')} | Submitted: {result.get('submitted')} | Rate: {result.get('response_rate')}%",
        f"Average: {result.get('average_score')} / 10 | Stars: {result.get('star_average')}",
        f"Generated: {datetime.utcnow().isoformat()}Z",
    ]
    for line in lines:
        c.drawString(40, y, line)
        y -= 16
    y -= 10
    c.setFont("Helvetica-Bold", 12)
    c.drawString(40, y, "Score distribution")
    y -= 18
    c.setFont("Helvetica", 10)
    for score, count in sorted((result.get("distribution") or {}).items()):
        c.drawString(50, y, f"Score {score}: {count}")
        y -= 14
        if y < 80:
            c.showPage()
            y = height - 50
    if result.get("ai_summary"):
        y -= 10
        c.setFont("Helvetica-Bold", 12)
        c.drawString(40, y, "Analysis summary")
        y -= 16
        c.setFont("Helvetica", 10)
        for chunk in _wrap(result["ai_summary"], 90):
            c.drawString(40, y, chunk)
            y -= 14
            if y < 80:
                c.showPage()
                y = height - 50
    c.save()
    return buf.getvalue()


def export_grade_pdf(campaign, result: dict[str, Any]) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    width, height = A4
    y = height - 50
    c.setFont("Helvetica-Bold", 14)
    c.drawString(40, y, "Grade / Ranking Report")
    y -= 24
    c.setFont("Helvetica", 11)
    for line in [
        f"Campaign: {campaign.name}",
        f"Group Head: {campaign.group_head.full_name if campaign.group_head else '-'}",
        f"Direction: {campaign.rank_direction.value}",
        f"Finalized: {result.get('is_final')}",
    ]:
        c.drawString(40, y, line)
        y -= 16
    y -= 8
    c.setFont("Helvetica-Bold", 12)
    c.drawString(40, y, "Final ranking")
    y -= 18
    c.setFont("Helvetica", 10)
    for row in result.get("ranking") or []:
        c.drawString(50, y, f"#{row.get('rank')} — {row.get('employee_name')}")
        y -= 14
        if y < 80:
            c.showPage()
            y = height - 50
    c.save()
    return buf.getvalue()


def _wrap(text: str, width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = ""
    for w in words:
        trial = f"{current} {w}".strip()
        if len(trial) <= width:
            current = trial
        else:
            if current:
                lines.append(current)
            current = w
    if current:
        lines.append(current)
    return lines


def build_ics(campaign) -> str:
    uid = f"campaign-{campaign.id}@leaderfeedback"
    dtstart = campaign.start_at.strftime("%Y%m%dT%H%M%SZ") if campaign.start_at else ""
    dtend = campaign.end_at.strftime("%Y%m%dT%H%M%SZ") if campaign.end_at else ""
    summary = f"Campaign: {campaign.name}"
    description = f"Mode: {campaign.mode.value}. Complete your assigned feedback/ranking."
    return (
        "BEGIN:VCALENDAR\r\n"
        "VERSION:2.0\r\n"
        "PRODID:-//LeaderFeedback//EN\r\n"
        "BEGIN:VEVENT\r\n"
        f"UID:{uid}\r\n"
        f"DTSTAMP:{datetime.utcnow().strftime('%Y%m%dT%H%M%SZ')}\r\n"
        f"DTSTART:{dtstart}\r\n"
        f"DTEND:{dtend}\r\n"
        f"SUMMARY:{summary}\r\n"
        f"DESCRIPTION:{description}\r\n"
        "END:VEVENT\r\n"
        "END:VCALENDAR\r\n"
    )
